"""PowerPoint extraction, slide citations, topic filtering and OCR of image-only pages."""

import io
import json

import pytest

from core import config, documents, ingest, ocr
from core.documents import DocumentError, Page


def make_pptx(slides: list[str], picture_on: int | None = None) -> bytes:
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    for i, text in enumerate(slides, 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
        if text:
            slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2)).text_frame.text = text
        if i == picture_on:
            image = io.BytesIO()
            Image.effect_noise((400, 300), 80).convert("RGB").save(image, "PNG")  # big enough to keep
            slide.shapes.add_picture(io.BytesIO(image.getvalue()), Inches(1), Inches(3))
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def image_pdf(pages: int) -> bytes:
    from PIL import Image

    out = io.BytesIO()
    frames = [Image.new("RGB", (600, 800), "white") for _ in range(pages)]
    frames[0].save(out, "PDF", save_all=True, append_images=frames[1:])
    return out.getvalue()


def test_pptx_slides_become_numbered_pages():
    data = make_pptx(["YARN splits resource management and scheduling.", "", "HDFS stores blocks."])
    pages = documents.extract_pages("bd.pptx", data)
    assert [(p.number, p.text) for p in pages] == [
        (1, "YARN splits resource management and scheduling."), (3, "HDFS stores blocks.")]
    assert len(documents.extract_pages("bd.pptx", data, keep_empty=True)) == 3


def test_slides_are_cited_as_slides():
    assert documents.location("Big Data - Unit 2.pptx", 4) == "Big Data - Unit 2.pptx, slide 4"
    assert documents.location("NLP - Unit 1.pdf", 4) == "NLP - Unit 1.pdf, p. 4"
    assert documents.location("notes.docx", None) == "notes.docx"


def test_image_only_file_without_ocr_is_rejected_clearly():
    with pytest.raises(DocumentError, match="No text found"):
        documents.extract_pages("scan.pdf", image_pdf(2))


def test_topics_only_listed_in_a_syllabus_line_are_dropped():
    reply = json.dumps({"topics": [
        {"name": "Hive tables", "coverage": "explained"},
        {"name": "Hive UDFs", "coverage": "listed_only"},
        {"name": "Sqoop import", "coverage": "explained"},
        {"name": "Sqoop export", "coverage": "brief"},
        {"name": "Pig macros", "coverage": "explained"},
    ]})
    assert ingest.parse_topics(reply) == ["Hive tables", "Sqoop import", "Pig macros"]
    # Short notes with few explained topics keep their brief ones; old plain lists still work.
    thin = json.dumps({"topics": [{"name": "A", "coverage": "explained"}, {"name": "B", "coverage": "brief"},
                                  {"name": "C", "coverage": "listed_only"}]})
    assert ingest.parse_topics(thin) == ["A", "B"]
    assert ingest.parse_topics('{"topics": ["X", "Y"]}') == ["X", "Y"]


@pytest.fixture
def fake_ocr(monkeypatch):
    monkeypatch.setattr(config, "OCR_ENABLED", True)
    monkeypatch.setattr(config, "ANTHROPIC_API_KEY", "test-key")
    calls = []

    def transcribe(content, what):
        calls.append(content)
        if what == "PDF pages":
            numbers = [int(n) for n in content[1]["text"].split("original pages ")[1].split(".")[0].split(", ")]
        else:
            numbers = [int(b["text"].split()[1].rstrip(":")) for b in content if b["type"] == "text"]
        return {n: f"Transcribed text of page {n} about context-free grammars." for n in numbers}, (100, 20)

    monkeypatch.setattr(ocr, "_transcribe", transcribe)
    return calls


def test_scanned_pdf_pages_are_read_in_batches_with_their_page_numbers(fake_ocr):
    pages = documents.extract_pages("nlp.pdf", data := image_pdf(8), keep_empty=True)
    result = ocr.fill_image_pages("nlp.pdf", data, pages)
    assert result.pages_read == 8 and result.input_tokens == 200
    assert len(fake_ocr) == 2  # 6 + 2 pages
    assert pages[7].number == 8 and "page 8" in pages[7].text


def test_only_picture_slides_are_sent(fake_ocr):
    data = make_pptx(["A full slide of real text " * 5, "", ""], picture_on=2)
    pages = documents.extract_pages("dl.pptx", data, keep_empty=True)
    ocr.fill_image_pages("dl.pptx", data, pages)
    assert len(fake_ocr) == 1 and "page 2" in pages[1].text
    assert pages[2].text == ""  # empty slide with no picture: nothing to read


def test_failed_batch_keeps_the_rest_of_the_file(fake_ocr, monkeypatch, tmp_path):
    from core import db, llm, vectorstore
    monkeypatch.setattr(ocr, "_transcribe", lambda content, what: None)
    monkeypatch.setattr(config, "DB_PATH", tmp_path / "t.db")
    monkeypatch.setattr(config, "UPLOADS_DIR", tmp_path / "uploads")
    monkeypatch.setattr(vectorstore, "add_chunks", lambda *a, **k: None)
    monkeypatch.setattr(llm, "generate", lambda *a, **k: '{"topics": ["Grammars"]}')
    db.init_db()
    uid = db.get_or_create_user("u")["id"]
    data = make_pptx(["Context-free grammars define constituents and phrase structure rules " * 3, ""],
                     picture_on=2)
    result = ingest.ingest(uid, "nlp.pptx", data)
    assert result.chunks == 1 and "1 image pages couldn't be read" in result.warning
